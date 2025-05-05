import streamlit as st
from langchain_groq import ChatGroq
import PyPDF2
from docx import Document as DocxDocument
from langchain.prompts import ChatPromptTemplate
from langchain_core.output_parsers import StrOutputParser
from fpdf import FPDF
import os
import sys
from dotenv import load_dotenv
import time

import spacy
from textblob import TextBlob
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

# Load .env
# load_dotenv()
# chat = ChatGroq(api_key=os.getenv('GROQ_API_KEY'), model_name="llama3-70b-8192")

if "GROQ_API_KEY" not in st.secrets:
    st.error("GROQ_API_KEY not found in Streamlit secrets")
else:
    chat = ChatGroq(api_key=st.secrets["GROQ_API_KEY"], model_name="llama3-70b-8192")

try:
    nlp = spacy.load("en_core_web_sm")
except OSError:
    print("Downloading spaCy model...")
    os.system(f"{sys.executable} -m spacy download en_core_web_sm")
    nlp = spacy.load("en_core_web_sm")

# Parse PDF
def parse_pdf(file):
    reader = PyPDF2.PdfReader(file)
    return "".join(page.extract_text() for page in reader.pages)

# Parse DOCX
def parse_docx(file):
    doc = DocxDocument(file)
    return "\n".join(para.text for para in doc.paragraphs)

# Extract Entities
def extract_entities(resume_text):
    doc = nlp(resume_text)
    return [(ent.text, ent.label_) for ent in doc.ents]

# Extract Keywords
def extract_keywords(text):
    vectorizer = TfidfVectorizer(stop_words='english')
    tfidf_matrix = vectorizer.fit_transform([text])
    feature_names = vectorizer.get_feature_names_out()
    sorted_items = tfidf_matrix.toarray().flatten().argsort()[::-1]
    return [feature_names[i] for i in sorted_items[:10]]

# Sentiment Analysis
def analyze_sentiment(answer):
    blob = TextBlob(answer)
    return blob.sentiment.polarity

# Similarity Check
def check_similarity(answers):
    vectorizer = TfidfVectorizer().fit_transform(answers)
    return cosine_similarity(vectorizer)

# LangChain Question Generator
@st.cache_data(ttl=600)
def fetch_questions(resume_text, num_qs,entities,keywords):
    prompt = ChatPromptTemplate.from_messages([
        ("system", "You are a helpful assistant."),
        ("human", f"""You are an experienced interviewer...
        Ask {num_qs} questions only. Each question must end with a £ except the last one. 
        \n\n{resume_text}\n\n Entitites : {entities} keywords : {keywords} Interview Questions:""")
    ])
    chain = prompt | chat | StrOutputParser()
    output = chain.invoke({"text": resume_text})
    return output.rstrip('£').split('£')

# Feedback Generator
def fetch_feedback(resume_text, combined_string):
    prompt = ChatPromptTemplate.from_messages([
        ("system", "You are a helpful assistant."),
        ("human", f"""Based on the following resume and answers, give honest feedback:
        \n\n{resume_text}\n\n{combined_string}\n\nFeedback:""")
    ])
    chain = prompt | chat | StrOutputParser()
    return chain.invoke({"text": resume_text})

# Evaluation Report Generator
def fetch_report(resume_text, combined_string):
    prompt = ChatPromptTemplate.from_messages([
        ("system", "You are a helpful assistant."),
        ("human", f"""Provide a detailed evaluation:
        Introduction to the candidate's performance
        Technical skills, Problem Solving, Communication, Teamwork
        Strengths, Areas of improvement, Overall impression, Recommendation.
        \n\n{resume_text}\n\n{combined_string}\n\nEvaluation:""")
    ])
    chain = prompt | chat | StrOutputParser()
    return chain.invoke({"text": resume_text})

# PDF Creator
def create_pdf(report_text):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=16)
    pdf.cell(200, 10, txt="Evaluation Report", ln=True, align='C')
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, report_text)

    local_directory = "report_pdfs"
    os.makedirs(local_directory, exist_ok=True)
    pdf_file_path = os.path.join(local_directory, "Evaluation.pdf")
    pdf.output(pdf_file_path)
    return pdf_file_path

# PDF Downloader
def download_pdf(file_path):
    with open(file_path, "rb") as file:
        return st.download_button(
            label="Download Report PDF",
            data=file,
            file_name="Evaluation.pdf",
            mime="application/pdf"
        )

# MAIN App
def main():
    st.title("AI Powered Mock Interview System")
    uploaded_file = st.file_uploader("Upload your resume (PDF/DOCX)", type=["pdf", "docx"])

    if uploaded_file:
        st.success("Resume uploaded successfully!")
        resume_text = parse_pdf(uploaded_file) if uploaded_file.type == "application/pdf" else parse_docx(uploaded_file)

        # NLP Analysis Section
        st.header("Resume Insights")
        entities = extract_entities(resume_text)
        keywords = extract_keywords(resume_text)

        # st.subheader("Named Entities")
        # st.write(entities)

        # st.subheader("Top Keywords")
        # st.write(keywords)

        num_qs = st.slider("Select number of interview questions:", 1, 10, 5)
        time.sleep(1)

        questions = fetch_questions(resume_text, num_qs,entities,keywords)
        responses = []
        st.header("💬 Interview Questions")
        for i, q in enumerate(questions):
            st.write(f"{i+1}. {q}")
            ans = st.text_area("Your Answer:", key=f"response_{i}", height=150)
            responses.append(ans)

        if st.button("Analyze Responses"):
            st.header("Sentiment of Your Answers")
            for i, ans in enumerate(responses):
                polarity = analyze_sentiment(ans)
                st.write(f"Q{i+1}: Sentiment Score: `{polarity}` {'👍' if polarity > 0 else '👎' if polarity < 0 else '😐'}")

            similarity_matrix = check_similarity(responses)
            st.subheader("Similarity Matrix")
            st.write(similarity_matrix)

        if st.button("Get Feedback"):
            combined = " ".join([f"{q}\n\nAnswer: {r}\n\n" for q, r in zip(questions, responses)])
            feedback = fetch_feedback(resume_text, combined)
            st.write("**Feedback:**")
            st.write(feedback)

        if st.button("📄 Generate PDF Report"):
            combined = " ".join([f"{q}\n\nAnswer: {r}\n\n" for q, r in zip(questions, responses)])
            evaluation = fetch_report(resume_text, combined)
            pdf_file_path = create_pdf(evaluation)
            st.success("✅ PDF generated successfully!")
            download_pdf(pdf_file_path)

if __name__ == "__main__":
    main()
